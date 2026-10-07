#!/usr/bin/env python3
"""Export ``benchmark.csv`` to one .xlsx for sharing: one sheet per dataset.

    uv run --with openpyxl python scripts/export_xlsx.py
    uv run --with openpyxl python scripts/export_xlsx.py --out results.xlsx

Each sheet opens with what the dataset is and how every method was run on it,
then lists all methods on that dataset -- DEPTHOR, DuCos, WAVE and the
no-training reference rows -- with every metric and, where the paper reports one,
the paper's RMSE next to ours in the paper's units.  The latest row per experiment
is used, so re-running after new experiments just refreshes the file.  Experiments
that exist as configs but have no row are listed on stdout.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import shutil
import tempfile
import time
from pathlib import Path

import _bootstrap  # noqa: F401
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.config import CONFIG_ROOT
from utils.results import RESULTS_DIR, migrate_summary_header

NEURAL = {"depthor": "DEPTHOR", "ducos": "DuCos", "wave": "WAVE"}
SHOWN = {**NEURAL, "nn_fill": "nn_fill", "local_bilateral": "local_bilateral", "bicubic": "bicubic"}

# experiment-name regex of every row that is exported.  wave_zju_l5_* and
# wave_void_stairs_* are invalid protocols (metrics_protocol.md) and stay out.
INCLUDE = (
    r"^(depthor|nn_fill|local_bilateral)_"
    r"|^(ducos_|bicubic_.*_x4$|bicubic_arkitscenes(_768)?$)"
    r"|^(wave_(minjiang|nyuv2|rgbd_stair|arkitscenes|hammer)|bicubic_.*_x(8|16|32)$)"
)

# experiment -> (paper RMSE, factor ours -> paper units, source)
PAPER: dict[str, tuple[float, float, str]] = {
    "depthor_zju_l5": (0.350, 1, "DEPTHOR, arXiv:2504.01596, Tab. 2 (Ours-Large), ZJU-L5, м"),
    "ducos_middlebury_x4": (1.45, 255, "DuCos, arXiv:2503.04171, Tab. 1, x4; шкала 0-255"),
    "ducos_lu_x4": (1.38, 255, "DuCos, arXiv:2503.04171, Tab. 1, x4; шкала 0-255"),
    "ducos_nyuv2_x4": (2.60, 100, "DuCos, arXiv:2503.04171, Tab. 1, x4; см"),
    "wave_nyuv2_x8": (2.50, 100, "WAVE, arXiv:2608.25302, Tab. 3 (обучен на NYU); см"),
    "wave_nyuv2_x16": (4.60, 100, "WAVE, arXiv:2608.25302, Tab. 3 (обучен на NYU); см"),
    "wave_nyuv2_x32": (7.90, 100, "WAVE, arXiv:2608.25302, Tab. 1 (обучен на NYU); см"),
}

DTOF = (
    "в каждой из 64 зон (сетка 8×8) берётся самая частая глубина, к ней добавляется шум "
    "(1 см + 1.2% расстояния), 5% зон выпадают; на вход идёт по одному пикселю на зону - "
    "так, как DEPTHOR обучали"
)
SR_NOTE = "веса не дообучались на этих данных (zero-shot)"

# sheet order: stairs first.  key = value of the `dataset` column in benchmark.csv
DATASETS: list[dict] = [
    {
        "key": "minjiang",
        "sheet": "MinJiang",
        "stairs": True,
        "about": (
            "Лестницы, RGB-D камера с двух точек зрения; берём сцену STAIRS, камеру CAM1: 890 кадров 640×480, "
            "глубина в метрах (uint16, мм). Отдельного эталона нет: глубина - сам датчик, поэтому все цифры "
            "оптимистичны, а дыры датчика (нули) не оцениваются."
        ),
        "runs": [
            f"DEPTHOR, nn_fill, local_bilateral: dToF 8×8 симулируется из эталона - {DTOF}.",
            "DuCos ×4, WAVE ×8/×16/×32 и bicubic: эталон (дыры заранее заполнены ближайшими значениями) уменьшается "
            f"бикубикой PIL в N раз, метод восстанавливает исходный размер; {SR_NOTE}.",
        ],
    },
    {
        "key": "rgbd_stair",
        "sheet": "RGB-D stair",
        "stairs": True,
        "about": (
            "Лестницы, тестовая часть набора StairNet (RGB-D камера): 154 кадра 640×480, из них 139 в оценке. "
            "Глубина записана в 8 бит с нормировкой по кадру; в метры переведена по диапазону из файла extrinsics "
            "(сверено с облаками точек), шаг квантования ≈17 мм. Отброшены 15 кадров с диапазоном глубины больше 10 м: "
            "при 256 уровнях шаг у них от 40 мм до 257 мм, эталон слишком грубый (в трёх кадрах диапазон 51-65 м, "
            "65535 - максимум uint16). Эталон - сам датчик, цифры оптимистичны."
        ),
        "runs": [
            f"DEPTHOR, nn_fill, local_bilateral: dToF 8×8 симулируется из эталона - {DTOF}.",
            f"DuCos ×4, WAVE ×8 и bicubic: эталон (дыры заполнены) уменьшается бикубикой PIL; {SR_NOTE}.",
        ],
    },
    {
        "key": "void_stairs",
        "sheet": "VOID",
        "stairs": True,
        "about": (
            "Лестницы (последовательности stairs0/1/3/4): 3459 кадров 640×480; в данных настоящие ~740 точек "
            "визуально-инерциальной одометрии и плотный эталон. Соседние кадры почти одинаковы."
        ),
        "runs": [
            "Строки с dToF: родные ~740 точек - не сетка 8×8 и DEPTHOR не подходят, поэтому dToF 8×8 симулируется из "
            f"плотного эталона ({DTOF}); цифры оптимистичны.",
            "Строка «~740 точек VIO» (local_bilateral): на родных точках датасета, без симуляции.",
            "DuCos и WAVE не запускались: 740 точек - слишком мало для задачи повышения разрешения "
            "(WAVE неотличим от bicubic).",
        ],
    },
    {
        "key": "zju_l5",
        "sheet": "ZJU-L5",
        "stairs": False,
        "about": (
            "Офисы, кафе, лаборатория; 527 тестовых кадров 640×480. Настоящий датчик dToF VL53L5CX 8×8 зон, эталон - "
            "стереореконструкция. Родной датасет DEPTHOR; не лестницы."
        ),
        "runs": [
            "DEPTHOR: вход - 64 значения настоящего датчика, по одному пикселю в центре каждой зоны (как в коде авторов), "
            "без симуляции. Это опорная точка: RMSE сверяется со статьёй.",
            "local_bilateral: те же 64 точки + Depth Anything V2 (без обучения под задачу). nn_fill, bicubic: "
            "ближайшая зона / бикубика по сетке зон, RGB не используется.",
            "DuCos и WAVE не запускались: 64 зоны - не вход для повышения разрешения.",
        ],
    },
    {
        "key": "hammer",
        "sheet": "HAMMER",
        "stairs": False,
        "about": (
            "Комнатные сцены с несколькими датчиками (тестовые сцены HAMMER): D435 (активное стерео, тот же принцип, что "
            "у Astra 2), L515 и Lucid Helios (ToF), эталон - лазер, независимый от датчиков. Кадры ≈1224×1024. Берётся "
            "каждый 5-й кадр каждой траектории (≈570 из 2828: соседние кадры почти одинаковы), одинаковые кадры во всех "
            "строках листа. Не лестницы."
        ),
        "runs": [
            f"Строки «симул. dToF 8×8 из <датчика>» (DEPTHOR, nn_fill, local_bilateral): 64 зоны строятся из глубины "
            f"этого датчика ({DTOF}). Настоящие дыры и тени датчика при этом теряются, остаются только 64 числа.",
            "Строки «настоящая карта <датчика>» (nn_fill, local_bilateral): на вход идёт сама карта датчика как есть, дыры "
            "заполняются ближайшими значениями. nn_fill - это «сам датчик»: насколько он ошибается против лазера. "
            "local_bilateral калибрует Depth Anything по плотным пикселям датчика (без обучения). DEPTHOR на такой плотной "
            "карте не запускали: его вход - 64 точки.",
            "DuCos ×4, WAVE ×8/×16/×32 и bicubic (только D435): карта датчика (дыры заранее заполнены) уменьшается "
            f"бикубикой PIL в N раз, метод восстанавливает исходный размер, оценка по лазеру; {SR_NOTE}.",
            "Столбцы «в дырах» и «вне дыр»: те же метрики отдельно по пикселям, где настоящий датчик ничего не измерил "
            "(в дырах - лазерный эталон есть, у датчика значения нет), и где измерил. Считаются по пикселям, не по кадрам. "
            "«Дыры датчика, %» - доля таких пикселей среди оцениваемых.",
            "Внизу листа - «разрыв доменов»: DEPTHOR на 64 зонах из D435, L515 и ToF на одних и тех же сценах.",
        ],
    },
    {
        "key": "arkitscenes",
        "sheet": "ARKitScenes",
        "stairs": False,
        "about": (
            "Комнаты; подмножество повышения разрешения, часть Validation: настоящий LiDAR iPad 256×192 на входе и "
            "лазерный эталон 1920×1440 (отношение 7.5). Число кадров - в столбце «Кадров» (весь набор ≈5600, при "
            "ускоренном прогоне каждый 10-й; кадры с повёрнутым RGB пропущены). Не лестницы."
        ),
        "runs": [
            f"DEPTHOR, nn_fill, local_bilateral: 64 зоны строятся из настоящей глубины LiDAR ({DTOF}), оценка по лазеру.",
            "WAVE и bicubic: вход - настоящая карта LiDAR без симуляции, кадр 1440×1920 (отношение 7.5); веса ×8 "
            "применяются к 7.5 - вне распределения.",
            "DuCos и свой bicubic: тот же LiDAR, но кадр уменьшен до 768×1024 - ровно ×4, как в обучении DuCos. На полном "
            "кадре DuCos не помещается в 48 ГБ памяти. Блоки с разным размером кадра между собой не сравнивать.",
        ],
    },
    {
        "key": "nyuv2",
        "sheet": "NYUv2",
        "stairs": False,
        "about": (
            "Комнаты; протокол статей по повышению разрешения: последние 449 из 1449 размеченных кадров "
            "(nyu_depth_v2_labeled.mat), 480×640, эталон - заполненная глубина Kinect. Не лестницы. Здесь есть "
            "числа из статей, столбцы «по статье» заполнены."
        ),
        "runs": [
            "DuCos ×4, WAVE ×8/×16/×32 и bicubic: эталон уменьшается бикубикой PIL в N раз (как у авторов), по краям "
            "обрезается 6 пикселей. В статьях RMSE дан в сантиметрах (наш × 100). DEPTHOR не запускался.",
        ],
    },
    {
        "key": "middlebury",
        "sheet": "Middlebury",
        "stairs": False,
        "about": (
            "Классический набор для повышения разрешения, 30 пар. Глубина - 8 бит, НЕ метры: все RMSE в нормированной "
            "шкале 0-1; число из статьи = наш RMSE × 255. Не лестницы."
        ),
        "runs": [
            "DuCos ×4 и bicubic: эталон уменьшается бикубикой PIL в 4 раза, обрезка 6 пикселей по краям."
        ],
    },
    {
        "key": "lu",
        "sheet": "Lu",
        "stairs": False,
        "about": (
            "Классический набор для повышения разрешения, 6 пар (ASUS Xtion). Глубина - 8 бит, НЕ метры: RMSE в шкале "
            "0-1; число из статьи = наш RMSE × 255. Не лестницы."
        ),
        "runs": [
            "DuCos ×4 и bicubic: эталон уменьшается бикубикой PIL в 4 раза, обрезка 6 пикселей по краям."
        ],
    },
]

# (header, csv field, number format); ↓ lower is better, ↑ higher
METRICS = [
    ("AbsRel ↓", "absrel", "0.0000"),
    ("RMSE ↓", "rmse", "0.0000"),
    ("RMSE(log) ↓", "rmse_log", "0.0000"),
    ("δ1 ↑", "delta1", "0.0000"),
    ("δ2 ↑", "delta2", "0.0000"),
    ("δ3 ↑", "delta3", "0.0000"),
    ("MAE ↓", "mae", "0.0000"),
    ("SqRel ↓", "sqrel", "0.0000"),
    ("SiLog ↓", "silog", "0.00"),
    ("FPS ↑", "fps", "0.0"),
    ("Задержка, мс ↓", "latency_ms_median", "0.0"),
]
# only on sheets whose rows carry them (eval.regions: sensor_holes): the same metrics split by
# where the REAL sensor had no measurement
REGION_METRICS = [
    ("Дыры датчика, % пикселей", "holes_frac", "0.0%"),
    ("AbsRel в дырах ↓", "holes_absrel", "0.0000"),
    ("RMSE в дырах ↓", "holes_rmse", "0.0000"),
    ("δ1 в дырах ↑", "holes_delta1", "0.0000"),
    ("AbsRel вне дыр ↓", "valid_absrel", "0.0000"),
    ("RMSE вне дыр ↓", "valid_rmse", "0.0000"),
    ("δ1 вне дыр ↑", "valid_delta1", "0.0000"),
]
PAPER_COLS = ("RMSE по статье", "× к ед. статьи", "Наш RMSE в ед. статьи", "Откл., %")
HEAD = (
    ["Эксперимент", "Вход (протокол)", "Метод", "Кадров"]
    + [m[0] for m in METRICS]
    + list(PAPER_COLS)
    + ["Дата"]
)
FONT = "Arial"
LEGEND = (
    "Голубым выделена нейросеть, остальные строки - ориентиры без обучения под задачу на том же входе. "
    "↓ - меньше лучше, ↑ - больше лучше; метрики считаются по пикселям, где у эталона есть значение. "
    "Сравнивать можно строки внутри одного блока (между линиями): нейросеть с ориентиром на том же входе. "
    "Блоки отличаются тем, сколько информации получает метод (dToF - 64 числа, ×32 - 300, ×8 - 4800), "
    "поэтому это не рейтинг методов между блоками; DuCos (×4) и WAVE (×8 и больше) тоже напрямую не сравнивать. "
    "FPS и задержка сняты на видеокартах RTX A6000 (кроме двух строк zju_l5/nn_fill и zju_l5/bicubic - они на CPU, "
    "их скорость не сравнивать) и зависят от размера кадра; DEPTHOR измерен через bpops_shim (замена CUDA-расширения на PyTorch), "
    "скорость с DuCos и WAVE не сравнивать."
)


def load_rows(path: Path) -> list[dict[str, str]]:
    """benchmark.csv under the current header, without touching the original."""
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / "benchmark.csv"
        shutil.copy(path, copy)
        migrate_summary_header(copy)
        with copy.open(newline="", encoding="utf-8") as fh:
            return list(csv.DictReader(fh))


def latest_per_experiment(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    """Newest row per experiment name; matrix runs ('adhoc') are keyed by dataset/method."""
    out: dict[str, dict[str, str]] = {}
    for r in sorted(rows, key=lambda r: r["timestamp"]):
        key = r["experiment"]
        if key == "adhoc":
            if r["dataset"] != "zju_l5" or r["method"] not in ("nn_fill", "bicubic"):
                continue  # the only matrix rows we keep: the floors on ZJU-L5, beside DEPTHOR
            key = f"{r['dataset']}/{r['method']}"
        out[key] = r
    return out


def hammer_sensor(exp: str) -> str:
    """Which HAMMER stream an experiment reads: the name carries it, D435 by default."""
    m = re.search(r"_(l515|tof)(?:_|$)", exp)
    return {"l515": "L515", "tof": "ToF"}[m.group(1)] if m else "D435"


def protocol(exp: str, r: dict[str, str]) -> tuple[int, str]:
    """(sort rank, label): the input the method was given."""
    ds, deg = r["dataset"], r["degradation"]
    if deg == "dtof_sim":
        src = (
            hammer_sensor(exp)
            if ds == "hammer"
            else "датчика"
            if ds == "arkitscenes"
            else "эталона"
        )
        return 0, f"симул. dToF 8×8 из {src}"
    m = re.search(r"_x(\d+)$", exp)
    if m:
        src = "карты D435" if ds == "hammer" else "эталона"
        return int(m.group(1)), f"bicubic ×{m.group(1)} из {src}"
    if ds == "hammer":
        return 0, f"настоящая карта {hammer_sensor(exp)} (с дырами)"
    if ds == "void_stairs":
        return 0, "~740 точек VIO"
    if ds == "arkitscenes":
        if exp.endswith("_768") or exp == "ducos_arkitscenes":
            return 101, "реальный LiDAR 256×192, кадр 768×1024 (×4)"
        return 100, "реальный LiDAR 256×192, кадр 1440×1920 (×7.5)"
    return 0, "реальный dToF 8×8"


def num(v: str):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def add_text(
    ws, row: int, text: str, last_col: int, *, bold=False, size=10, italic=False, chars=150
) -> int:
    """One wrapped paragraph merged across the table width; returns the next free row."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, bold=bold, size=size, italic=italic)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 14 * max(1, math.ceil(len(text) / chars)) + 2
    return row + 1


def build(rows: list[dict[str, str]], configs: set[str]) -> tuple[Workbook, list[str]]:
    latest = latest_per_experiment(rows)
    # "dataset/method" keys are the matrix floors on ZJU-L5
    picked = {e: r for e, r in latest.items() if re.search(INCLUDE, e) or "/" in e}
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="999999")

    for d in DATASETS:
        ws = wb.create_sheet(d["sheet"])
        if d["stairs"]:
            ws.sheet_properties.tabColor = "ED7D31"

        items = []
        for exp, r in picked.items():
            if r["dataset"] != d["key"] or r["method"] not in SHOWN:
                continue
            rank, label = protocol(exp, r)
            items.append((rank, label, r["method"] not in NEURAL, r["method"], exp, r))
        items.sort(key=lambda t: t[:5])

        regions = any(num(r.get("holes_rmse", "")) is not None for *_, r in items)
        head = list(HEAD)
        if regions:
            at = head.index(PAPER_COLS[0])
            head[at:at] = [m[0] for m in REGION_METRICS]
        head_cols = len(head)

        title = f"{d['sheet']}: лестницы" if d["stairs"] else d["sheet"]
        r_i = add_text(ws, 1, title, head_cols, bold=True, size=13)
        r_i = add_text(ws, r_i, "Что это. " + d["about"], head_cols)
        r_i = add_text(ws, r_i, "Как запускали методы:", head_cols, bold=True)
        for line in d["runs"]:
            r_i = add_text(ws, r_i, "• " + line, head_cols)
        r_i = add_text(ws, r_i, LEGEND, head_cols, size=9, italic=True)
        head_row = r_i + 1

        for j, h in enumerate(head, 1):
            c = ws.cell(head_row, j, h)
            c.font = Font(name=FONT, bold=True, size=10)
            c.fill = PatternFill("solid", fgColor="D9D9D9")
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            c.border = Border(bottom=Side(style="medium"))
        col = {h: j for j, h in enumerate(head, 1)}

        prev_label, row_i = None, head_row + 1
        at_row: dict[str, int] = {}  # experiment -> sheet row, for the domain-gap block
        for _, label, floor, method, exp, r in items:
            at_row[exp] = row_i
            fill = None if floor else PatternFill("solid", fgColor="DDEBF7")
            vals = [exp, label, SHOWN[method], num(r["n_samples"])]
            vals += [num(r[m[1]]) for m in METRICS]
            if regions:
                vals += [num(r.get(m[1], "")) for m in REGION_METRICS]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row_i, j, v)
                c.font = Font(name=FONT, size=10, bold=(not floor) and j in (1, 3))
            for name, _, fmt in METRICS:
                ws.cell(row_i, col[name]).number_format = fmt
            if regions:
                for name, _, fmt in REGION_METRICS:
                    ws.cell(row_i, col[name]).number_format = fmt
            ws.cell(row_i, col["Кадров"]).number_format = "#,##0"

            paper = PAPER.get(exp)
            if paper:
                p_val, factor, src = paper
                rmse_cell = f"{get_column_letter(col['RMSE ↓'])}{row_i}"
                pc, fc, oc, dc = (col[n] for n in PAPER_COLS)
                ws.cell(row_i, pc, p_val).font = Font(name=FONT, size=10, color="0000FF")
                ws.cell(row_i, pc).comment = Comment(src, "export_xlsx")
                ws.cell(row_i, fc, factor).font = Font(name=FONT, size=10, color="0000FF")
                p_ref, f_ref = f"{get_column_letter(pc)}{row_i}", f"{get_column_letter(fc)}{row_i}"
                ws.cell(row_i, oc, f"={rmse_cell}*{f_ref}").number_format = "0.000"
                ws.cell(
                    row_i, dc, f"=({get_column_letter(oc)}{row_i}-{p_ref})/{p_ref}"
                ).number_format = "0.0%"
                for n in ("Наш RMSE в ед. статьи", "Откл., %"):
                    ws.cell(row_i, col[n]).font = Font(name=FONT, size=10)

            stamp = r["timestamp"]
            day = f"{stamp[:4]}-{stamp[4:6]}-{stamp[6:8]}"
            ws.cell(row_i, col["Дата"], day).font = Font(name=FONT, size=9)
            for j in range(1, head_cols + 1):
                c = ws.cell(row_i, j)
                if fill:
                    c.fill = fill
                if prev_label is not None and label != prev_label:
                    c.border = Border(top=thin)
            prev_label, row_i = label, row_i + 1

        if row_i == head_row + 1:
            ws.cell(row_i, 1, "результатов пока нет").font = Font(name=FONT, size=10, italic=True)
        widths = {
            "Эксперимент": 30,
            "Вход (протокол)": 28,
            "Метод": 16,
            "Кадров": 8,
            "Дата": 11,
        }
        for h, j in col.items():
            ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 11)
        ws.row_dimensions[head_row].height = 42
        ws.freeze_panes = ws.cell(head_row + 1, 3)
        ws.auto_filter.ref = (
            f"A{head_row}:{get_column_letter(head_cols)}{max(row_i - 1, head_row + 1)}"
        )

        if d["key"] == "hammer":
            domain_gap_block(ws, picked, row_i + 2)

    missing = sorted(e for e in configs if re.search(INCLUDE, e) and e not in latest)
    return wb, missing


GAP_ROWS = (
    ("D435 (активное стерео)", "depthor_hammer_dtof", "nn_fill_hammer_d435"),
    ("L515", "depthor_hammer_l515_dtof", "nn_fill_hammer_l515"),
    ("Lucid Helios (I-ToF)", "depthor_hammer_tof_dtof", "nn_fill_hammer_tof"),
)


def domain_gap_block(ws, picked: dict[str, dict[str, str]], top: int) -> None:
    """DEPTHOR fed 64 zones from each HAMMER sensor, same scenes and laser GT.

    The gap is the change against the D435 row (the sensor closest to Astra 2).  Raw
    sensor columns say how wrong each sensor is before any model touches it.
    """
    if not any(e in picked for _, e, _ in GAP_ROWS):
        return
    heads = [
        "Разрыв доменов: DEPTHOR на 64 зонах разных датчиков",
        "AbsRel ↓",
        "RMSE, м ↓",
        "δ1 ↑",
        "AbsRel к D435, %",
        "RMSE к D435, %",
        "Сам датчик: AbsRel ↓",
        "Сам датчик: RMSE, м ↓",
    ]
    for j, h in enumerate(heads, 1):
        c = ws.cell(top, j, h)
        c.font = Font(name=FONT, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[top].height = 42
    ref = top + 1  # the D435 row
    for k, (name, depthor, raw) in enumerate(GAP_ROWS):
        r_i = top + 1 + k
        a, b = picked.get(depthor), picked.get(raw)
        vals = [
            name,
            num(a["absrel"]) if a else None,
            num(a["rmse"]) if a else None,
            num(a["delta1"]) if a else None,
        ]
        for j, v in enumerate(vals, 1):
            ws.cell(r_i, j, v).font = Font(name=FONT, size=10)
        if a and k:
            ws.cell(r_i, 5, f"=(B{r_i}-B${ref})/B${ref}").number_format = "0.0%"
            ws.cell(r_i, 6, f"=(C{r_i}-C${ref})/C${ref}").number_format = "0.0%"
        ws.cell(r_i, 7, num(b["absrel"]) if b else None)
        ws.cell(r_i, 8, num(b["rmse"]) if b else None)
        for j in (2, 3, 4, 7, 8):
            ws.cell(r_i, j).number_format = "0.0000"
        for j in range(5, 9):
            ws.cell(r_i, j).font = Font(name=FONT, size=10)
    note = (
        "Одни и те же сцены и лазерный эталон; меняется только датчик, из которого строятся 64 зоны. "
        "Разрыв - изменение ошибки относительно D435; D435 ближе всего к Astra 2."
    )
    ws.cell(top + 1 + len(GAP_ROWS), 1, note).font = Font(name=FONT, size=9, italic=True)


def main() -> int:
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    p.add_argument("--csv", type=Path, default=RESULTS_DIR / "benchmark.csv")
    p.add_argument(
        "--out", type=Path, default=RESULTS_DIR / f"baselines_{time.strftime('%Y-%m-%d')}.xlsx"
    )
    args = p.parse_args()

    rows = load_rows(args.csv)
    configs = {
        q.stem for q in (CONFIG_ROOT / "experiment").glob("*.yaml") if not q.stem.startswith("_")
    }
    wb, missing = build(rows, configs)
    wb.save(args.out)
    print(f"{len(rows)} rows read, wrote {args.out}")
    if missing:
        print(f"no result yet for: {' '.join(missing)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
